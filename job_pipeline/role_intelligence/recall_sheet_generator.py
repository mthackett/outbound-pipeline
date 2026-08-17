#!/usr/bin/env python3
"""Configurable Resume Recall Sheet / Candidate Cockpit DOCX generator.

The renderer is content-agnostic. It supports one or more landscape pages,
3-column layouts, runtime variable substitution, optional generation manifests,
and the original one-page content format for backward compatibility.

CLI
---
python recall_sheet_generator.py \
  --content example_content.json \
  --output candidate_cockpit.docx \
  --set posted_pay='$110K-$130K' \
  --set target_pay='$125K'

Python
------
from recall_sheet_generator import generate_recall_sheet, load_json

content = load_json("example_content.json")
generate_recall_sheet(
    content,
    "candidate_cockpit.docx",
    variables={"posted_pay": "$110K-$130K"},
)
"""

from __future__ import annotations

import argparse
import copy
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PACKAGE_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG_PATH = PACKAGE_DIR / "default_config.json"
VAR_PATTERN = re.compile(r"\{\{\s*([A-Za-z0-9_.-]+)\s*\}\}")
TOOL_VERSION = "1.0.0"


def load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def deep_merge(base: dict[str, Any], override: Mapping[str, Any]) -> dict[str, Any]:
    """Recursively merge override values into a copy of base."""
    out = copy.deepcopy(base)
    for key, value in override.items():
        if isinstance(value, Mapping) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = copy.deepcopy(value)
    return out


def load_config(config: str | Path | Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Load bundled defaults, optionally applying overrides."""
    base = load_json(DEFAULT_CONFIG_PATH)
    if config is None:
        return base
    if isinstance(config, (str, Path)):
        return deep_merge(base, load_json(config))
    return deep_merge(base, config)


def _pages_from_content(content: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    """Return modern pages, while accepting the original one-page structure."""
    if isinstance(content.get("pages"), list) and content["pages"]:
        return list(content["pages"])

    if "title" in content and "columns" in content:
        pages: list[Mapping[str, Any]] = [
            {"title": content["title"], "columns": content["columns"]}
        ]
        if isinstance(content.get("page2"), Mapping):
            pages.append(content["page2"])
        return pages

    raise ValueError("content must contain either pages[] or legacy title + columns")


def validate_content(content: Mapping[str, Any]) -> None:
    """Dependency-free structural validation with useful error messages."""
    pages = _pages_from_content(content)
    allowed_kinds = {"bullets", "lines", "skills", "mixed"}

    for pi, page in enumerate(pages):
        if not isinstance(page.get("title"), str) or not page["title"].strip():
            raise ValueError(f"pages[{pi}].title must be a non-empty string")
        columns = page.get("columns")
        if not isinstance(columns, list) or not columns:
            raise ValueError(f"pages[{pi}].columns must be a non-empty list")
        for ci, column in enumerate(columns):
            if not isinstance(column, list):
                raise ValueError(f"pages[{pi}].columns[{ci}] must be a list of sections")
            for si, section in enumerate(column):
                if not isinstance(section, Mapping):
                    raise ValueError(
                        f"pages[{pi}].columns[{ci}][{si}] must be an object"
                    )
                if not str(section.get("title", "")).strip():
                    raise ValueError(
                        f"pages[{pi}].columns[{ci}][{si}].title must be non-empty"
                    )
                kind = section.get("kind", "bullets")
                if kind not in allowed_kinds:
                    raise ValueError(
                        f"pages[{pi}].columns[{ci}][{si}].kind={kind!r} is unsupported; "
                        f"expected one of {sorted(allowed_kinds)}"
                    )


def _resolve_variables(
    obj: Any,
    variables: Mapping[str, Any],
    *,
    unset_text: str,
) -> Any:
    """Recursively replace {{variable}} tokens in strings."""
    if isinstance(obj, str):
        def repl(match: re.Match[str]) -> str:
            key = match.group(1)
            value = variables.get(key)
            if value is None or str(value).strip() == "":
                return unset_text
            return str(value)
        return VAR_PATTERN.sub(repl, obj)
    if isinstance(obj, list):
        return [_resolve_variables(v, variables, unset_text=unset_text) for v in obj]
    if isinstance(obj, Mapping):
        return {
            k: _resolve_variables(v, variables, unset_text=unset_text)
            for k, v in obj.items()
        }
    return obj


def _hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.strip().lstrip("#")
    if len(value) != 6:
        raise ValueError(f"Expected 6-digit hex color, got {value!r}")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def _set_font(run, cfg, *, size=None, bold=False, color=None, italic=False):
    name = cfg["font_name"]
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size if size is not None else cfg["body_font_pt"])
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*_hex_rgb(color))


def _set_para_spacing(p, *, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _set_cell_margins(cell, top=0, start=70, bottom=0, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top), ("start", start), ("bottom", bottom), ("end", end)
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def _set_table_borders_none(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def _set_table_borders(table, color="D9E2F3", size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color.lstrip("#"))


def _set_repeat_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _clear_initial_paragraph(cell):
    p = cell.paragraphs[0]
    _set_para_spacing(p, after=0, line=0.1)
    r = p.add_run("")
    r.font.size = Pt(1)


def _add_section_banner(cell, title, cfg):
    psp = cfg["spacing"]
    outer = cell.add_table(rows=1, cols=1)
    outer.autofit = False
    _set_table_borders_none(outer)
    banner_cell = outer.cell(0, 0)
    _set_cell_shading(banner_cell, cfg["banner_color"])
    _set_cell_margins(
        banner_cell,
        top=cfg["banner_pad_twips_top"],
        bottom=cfg["banner_pad_twips_bottom"],
        start=cfg["banner_pad_twips_side"],
        end=cfg["banner_pad_twips_side"],
    )
    p = banner_cell.paragraphs[0]
    _set_para_spacing(p, before=0, after=0, line=1.0)
    r = p.add_run(title.upper())
    _set_font(
        r,
        cfg,
        size=cfg["banner_font_pt"],
        bold=True,
        color=cfg["banner_text_color"],
    )

    spacer = cell.add_paragraph()
    _set_para_spacing(spacer, after=psp["after_banner_pt"], line=0.1)
    spacer.add_run("").font.size = Pt(1)


def _add_bullet(cell, item, cfg):
    p = cell.add_paragraph()
    _set_para_spacing(
        p,
        after=cfg["spacing"]["bullet_after_pt"],
        line=cfg["spacing"]["body_line"],
    )
    p.paragraph_format.left_indent = Pt(cfg["bullet_left_pt"])
    p.paragraph_format.first_line_indent = Pt(-cfg["bullet_hanging_pt"])
    bullet = p.add_run("• ")
    _set_font(bullet, cfg, size=cfg["body_font_pt"], color=cfg["body_color"])
    lead = item.get("lead", "")
    rest = item.get("rest", "")
    if lead:
        r = p.add_run(str(lead))
        _set_font(
            r,
            cfg,
            size=cfg["body_font_pt"],
            bold=item.get("lead_bold", True),
            color=cfg["body_color"],
        )
    if rest:
        r = p.add_run(str(rest))
        _set_font(r, cfg, size=cfg["body_font_pt"], color=cfg["body_color"])
    return p


def _add_line(cell, item, cfg):
    p = cell.add_paragraph()
    _set_para_spacing(
        p,
        after=cfg["spacing"]["line_after_pt"],
        line=cfg["spacing"]["body_line"],
    )
    label = item.get("label", "")
    value = item.get("value", "")
    if label:
        r = p.add_run(str(label))
        _set_font(r, cfg, size=cfg["body_font_pt"], bold=True, color=cfg["accent_color"])
    r = p.add_run(str(value))
    _set_font(r, cfg, size=cfg["body_font_pt"], color=cfg["body_color"])
    return p


def _add_skill_matrix(cell, items, cfg):
    if not items:
        return
    table = cell.add_table(rows=len(items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_borders(table, cfg["skill_border_color"], size=cfg["skill_border_size"])
    label_w = Inches(cfg["skill_label_width_in"])
    text_w = Inches(cfg["skill_text_width_in"])
    table.columns[0].width = label_w
    table.columns[1].width = text_w

    grid_cols = table._tbl.tblGrid.gridCol_lst
    if len(grid_cols) >= 2:
        grid_cols[0].set(qn("w:w"), str(int(cfg["skill_label_width_in"] * 1440)))
        grid_cols[1].set(qn("w:w"), str(int(cfg["skill_text_width_in"] * 1440)))

    for row, item in zip(table.rows, items):
        _set_repeat_no_split(row)
        label_cell, text_cell = row.cells
        label_cell.width = label_w
        text_cell.width = text_w
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(label_cell, top=45, bottom=45, start=55, end=55)
        _set_cell_margins(text_cell, top=45, bottom=45, start=70, end=55)
        _set_cell_shading(label_cell, cfg["skill_label_fill"])

        p = label_cell.paragraphs[0]
        _set_para_spacing(p, after=0, line=1.0)
        r = p.add_run(str(item.get("label", "")))
        _set_font(
            r,
            cfg,
            size=cfg["skill_label_font_pt"],
            bold=True,
            color=cfg["accent_color"],
        )

        p = text_cell.paragraphs[0]
        _set_para_spacing(p, after=0, line=cfg["spacing"]["skill_line"])
        terms = item.get("terms", [])
        text = "  •  ".join(map(str, terms)) if isinstance(terms, list) else str(terms)
        r = p.add_run(text)
        _set_font(r, cfg, size=cfg["skill_font_pt"], color=cfg["body_color"])

    spacer = cell.add_paragraph()
    _set_para_spacing(spacer, after=cfg["spacing"]["skill_after_pt"], line=0.1)
    spacer.add_run("").font.size = Pt(1)


def _add_note(cell, text, cfg):
    p = cell.add_paragraph()
    _set_para_spacing(
        p,
        after=cfg["spacing"]["note_after_pt"],
        line=cfg["spacing"]["note_line"],
    )
    r = p.add_run(str(text))
    _set_font(r, cfg, size=cfg["note_font_pt"], color=cfg["note_color"])


def _render_section(cell, section, cfg):
    _add_section_banner(cell, section["title"], cfg)
    kind = section.get("kind", "bullets")

    if kind == "bullets":
        for item in section.get("items", []):
            _add_bullet(cell, item, cfg)
    elif kind == "lines":
        for item in section.get("items", []):
            _add_line(cell, item, cfg)
    elif kind == "skills":
        _add_skill_matrix(cell, section.get("items", []), cfg)
    elif kind == "mixed":
        for block in section.get("blocks", []):
            bkind = block.get("kind")
            if bkind == "bullet":
                _add_bullet(cell, block, cfg)
            elif bkind == "line":
                _add_line(cell, block, cfg)
            elif bkind == "note":
                _add_note(cell, block.get("text", ""), cfg)
            else:
                raise ValueError(f"Unsupported mixed block kind: {bkind}")

    if section.get("note"):
        _add_note(cell, section["note"], cfg)


def _render_page(doc: Document, page: Mapping[str, Any], cfg: Mapping[str, Any]) -> None:
    columns = page["columns"]
    if len(columns) != len(cfg["column_widths_in"]):
        raise ValueError(
            "Number of content columns must match config column_widths_in "
            f"({len(columns)} content columns vs {len(cfg['column_widths_in'])} configured widths)"
        )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_para_spacing(p, after=cfg["spacing"]["title_after_pt"], line=1.0)
    r = p.add_run(str(page["title"]))
    _set_font(r, cfg, size=cfg["title_font_pt"], bold=True, color=cfg["accent_color"])

    body = doc.add_table(rows=1, cols=len(columns))
    body.alignment = WD_TABLE_ALIGNMENT.CENTER
    body.autofit = False
    _set_table_borders_none(body)

    for cell, width, sections in zip(body.rows[0].cells, cfg["column_widths_in"], columns):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(
            cell,
            top=0,
            start=cfg["column_pad_twips"],
            bottom=0,
            end=cfg["column_pad_twips"],
        )
        _clear_initial_paragraph(cell)
        for section in sections:
            _render_section(cell, section, cfg)

    _set_repeat_no_split(body.rows[0])


def generate_recall_sheet(
    content: Mapping[str, Any],
    output_path: str | Path,
    *,
    config: str | Path | Mapping[str, Any] | None = None,
    variables: Mapping[str, Any] | None = None,
    manifest_path: str | Path | None = None,
) -> Path:
    """Generate a one- or multi-page candidate cockpit DOCX.

    If ``manifest_path`` is supplied, also write a JSON sidecar recording the
    job/resume identifiers, tool/template version, variables, and output path.
    Use ``manifest_path="auto"`` to write ``<output>.manifest.json``.
    """
    validate_content(content)
    cfg = load_config(config)

    merged_variables = dict(content.get("variables", {}))
    if variables:
        merged_variables.update(variables)
    resolved = _resolve_variables(
        content,
        merged_variables,
        unset_text=cfg.get("unset_variable_text", "Set per role"),
    )
    pages = _pages_from_content(resolved)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(cfg["page"]["width_in"])
    sec.page_height = Inches(cfg["page"]["height_in"])
    sec.top_margin = Inches(cfg["page"]["top_margin_in"])
    sec.bottom_margin = Inches(cfg["page"]["bottom_margin_in"])
    sec.left_margin = Inches(cfg["page"]["left_margin_in"])
    sec.right_margin = Inches(cfg["page"]["right_margin_in"])

    normal = doc.styles["Normal"]
    normal.font.name = cfg["font_name"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), cfg["font_name"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), cfg["font_name"])
    normal.font.size = Pt(cfg["body_font_pt"])
    normal.paragraph_format.space_after = Pt(0)

    for idx, page in enumerate(pages):
        if idx:
            doc.add_page_break()
        _render_page(doc, page, cfg)

    props = doc.core_properties
    first_title = str(pages[0]["title"])
    props.title = resolved.get("document_title", first_title)
    props.subject = resolved.get("subject", "Resume recall and interview playbook")
    props.author = resolved.get("author", "")

    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        raise ValueError("output_path must end in .docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    if manifest_path is not None:
        if str(manifest_path).lower() == "auto":
            manifest_file = output_path.with_suffix(".manifest.json")
        else:
            manifest_file = Path(manifest_path)
        manifest_file.parent.mkdir(parents=True, exist_ok=True)
        manifest = {
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "tool": "candidate-cockpit",
            "tool_version": TOOL_VERSION,
            "template_version": resolved.get("template_version", "1.0.0"),
            "output_docx": str(output_path),
            "job": resolved.get("job", {}),
            "resume": resolved.get("resume", {}),
            "generation": resolved.get("generation", {}),
            "variables": merged_variables,
        }
        manifest_file.write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    return output_path


def generate_from_files(
    content_path: str | Path,
    output_path: str | Path,
    *,
    config_path: str | Path | None = None,
    variables: Mapping[str, Any] | None = None,
    manifest_path: str | Path | None = None,
) -> Path:
    content = load_json(content_path)
    return generate_recall_sheet(
        content,
        output_path,
        config=config_path,
        variables=variables,
        manifest_path=manifest_path,
    )


def _parse_cli_vars(values: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for value in values:
        if "=" not in value:
            raise ValueError(f"--set expects KEY=VALUE, got {value!r}")
        key, val = value.split("=", 1)
        key = key.strip()
        if not key:
            raise ValueError("--set key cannot be empty")
        out[key] = val
    return out


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate a configurable resume recall sheet / candidate cockpit DOCX"
    )
    parser.add_argument("--content", required=True, help="Path to cockpit content JSON")
    parser.add_argument(
        "--config",
        default=str(DEFAULT_CONFIG_PATH),
        help="Path to style/layout config JSON",
    )
    parser.add_argument("--output", required=True, help="Output .docx path")
    parser.add_argument(
        "--set",
        action="append",
        default=[],
        metavar="KEY=VALUE",
        help="Override a {{variable}} token. Repeat as needed.",
    )
    parser.add_argument(
        "--manifest",
        default=None,
        help="Optional manifest path, or 'auto' for <output>.manifest.json",
    )
    args = parser.parse_args()

    runtime_vars = _parse_cli_vars(args.set)
    output = generate_from_files(
        args.content,
        args.output,
        config_path=args.config,
        variables=runtime_vars,
        manifest_path=args.manifest,
    )
    print(output)


if __name__ == "__main__":
    main()
